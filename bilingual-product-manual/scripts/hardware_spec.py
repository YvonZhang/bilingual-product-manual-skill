"""Native editable specification pages with source-derived typography and tables."""
from datetime import datetime,timezone
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree
from PIL import Image
from docx import Document
from docx.shared import Pt,Emu,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT,WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from hardware_common import ROOT,LABELS,text,require,asset_path
from hardware_brochure import height

def el(tag,**attributes):
    node=OxmlElement('w:'+tag)
    for k,v in attributes.items():node.set(qn('w:'+k),str(v))
    return node

def font(run,t,size,bold=False,color='ink'):
    run.font.name=t['font_family'];run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(t['colors'][color]);rp=run._element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=el('rFonts');rp.insert(0,rf)
    for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),t['font_family'])

def position_logo(shape):
    """Anchor to the shared page origin, avoiding renderer-specific inline padding."""
    inline=shape._inline;anchor=OxmlElement('wp:anchor')
    for k,v in {'distT':0,'distB':0,'distL':0,'distR':0,'simplePos':0,'relativeHeight':1,'behindDoc':0,'locked':0,'layoutInCell':1,'allowOverlap':1}.items():anchor.set(k,str(v))
    simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');anchor.append(simple)
    for axis in ['H','V']:
        pos=OxmlElement('wp:position'+axis);pos.set('relativeFrom','page');offset=OxmlElement('wp:posOffset');offset.text=str(int(Pt(36)));pos.append(offset);anchor.append(pos)
    anchor.append(inline.find(qn('wp:extent')))
    effect=OxmlElement('wp:effectExtent')
    for side in ['l','t','r','b']:effect.set(side,'0')
    anchor.append(effect);anchor.append(OxmlElement('wp:wrapNone'))
    for child in list(inline):anchor.append(child)
    inline.getparent().replace(inline,anchor)

def configure(doc,t):
    s=doc.sections[0];v=t['specification'];s.page_width=Emu(v['page_width_twips']*635);s.page_height=Emu(v['page_height_twips']*635)
    for name,key in [('top_margin','margin_top_twips'),('bottom_margin','margin_bottom_twips'),('left_margin','margin_left_twips'),('right_margin','margin_right_twips'),('header_distance','header_twips'),('footer_distance','footer_twips')]:setattr(s,name,Emu(v[key]*635))
    s.different_first_page_header_footer=False
    for old in list(s._sectPr):
        if old.tag==qn('w:pgBorders'):s._sectPr.remove(old)
    border=el('pgBorders',offsetFrom='page');border.append(el('top',val='single',sz=32,color=t['colors']['accent'],space=0));s._sectPr.append(border)
    for name,size,leading,after,before in [('Normal',10.5,15.55,5,0),('Title',27,39.95,7,0),('Heading 1',23,34.05,9,0),('Heading 2',14,20.7,9,10),('Header',8.5,12,0,0),('Footer',7.5,11,0,0)]:
        st=doc.styles[name];st.font.name=t['font_family'];st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(t['colors']['ink']);st.font.bold=name in ['Title','Heading 1','Heading 2']
        for old in list(st.element.get_or_add_rPr().findall(qn('w:rFonts'))):st.element.get_or_add_rPr().remove(old)
        st.element.get_or_add_rPr().append(el('rFonts',ascii=t['font_family'],hAnsi=t['font_family'],eastAsia=t['font_family'],cs=t['font_family']))
        p=st.paragraph_format;p.space_before=Pt(before);p.space_after=Pt(after);p.line_spacing=Pt(leading);p.keep_with_next=name in ['Title','Heading 1','Heading 2'];p.widow_control=True
        p.left_indent=Pt(0);p.right_indent=Pt(0);p.first_line_indent=Pt(0)
        if name in ['Header','Footer']:p.tab_stops.clear_all()
        pr=st.element.get_or_add_pPr()
        for old in list(pr.findall(qn('w:snapToGrid'))):pr.remove(old)
        pr.append(el('snapToGrid',val='false'))
        for tag in ['pBdr','shd']:
            x=st.element.get_or_add_pPr().find(qn('w:'+tag))
            if x is not None:st.element.get_or_add_pPr().remove(x)
    return s

def create_template(path,t):
    doc=Document();configure(doc,t)
    for p in list(doc.paragraphs):p._element.getparent().remove(p._element)
    for k in ['author','last_modified_by','title','subject','keywords','comments','category','identifier']:setattr(doc.core_properties,k,'')
    buffer=BytesIO();doc.save(buffer);buffer.seek(0)
    # python-docx's starter carries empty bibliography custom XML and a stock
    # preview. Neither is part of this template's reusable style contract.
    with ZipFile(buffer) as source,ZipFile(path,'w',ZIP_DEFLATED) as target:
        for name in source.namelist():
            if name.startswith('customXml/') or name=='docProps/thumbnail.jpeg':continue
            content=source.read(name)
            if name.endswith('.rels') or name=='[Content_Types].xml':
                tree=etree.fromstring(content)
                for child in list(tree):
                    destination=child.get('Target','')+child.get('PartName','')
                    if 'customXml/' in destination or 'thumbnail.jpeg' in destination:tree.remove(child)
                content=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
            target.writestr(name,content)

class Spec:
    def __init__(self,data,t,lang,base,logo):
        self.data=data;self.t=t;self.lang=lang;self.base=base;self.logo=logo;self.doc=Document(ROOT/'assets/hardware-specification.docx');self.section=configure(self.doc,t);self.width=t['specification']['content_width_twips']/20;self.expected=[];self.layout=[]
        self.frame()

    def add(self,s,style='Normal',size=None,leading=None,bold=False,color='ink',after=None,before=None,page_break=False,align=None,keep=False,record=True):
        p=self.doc.add_paragraph(style=style);p.paragraph_format.page_break_before=page_break;p.paragraph_format.keep_with_next=keep or style in ['Title','Heading 1','Heading 2'];p.paragraph_format.keep_together=True
        p.paragraph_format.widow_control=True
        if after is not None:p.paragraph_format.space_after=Pt(after)
        if before is not None:p.paragraph_format.space_before=Pt(before)
        if leading is not None:p.paragraph_format.line_spacing=Pt(leading)
        if align is not None:p.alignment=align
        r=p.add_run(s)
        if size is not None:font(r,self.t,size,bold,color)
        else:
            r.font.name=self.t['font_family'];r._element.get_or_add_rPr().append(el('rFonts',ascii=self.t['font_family'],hAnsi=self.t['font_family'],eastAsia=self.t['font_family']))
        if record and s:self.expected.append(s)
        return p

    def frame(self):
        s=self.section;header=s.header
        for p in list(header.paragraphs):p._element.getparent().remove(p._element)
        p=header.add_paragraph(style='Header');p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(29);p.paragraph_format.tab_stops.add_tab_stop(Pt(self.width),WD_TAB_ALIGNMENT.RIGHT)
        if self.logo:
            im=Image.open(self.logo);w=self.t['logo_visible_width_mm']*72/25.4;h=w*im.height/im.width;require(h<=29,'Logo is too tall at its shared visual width')
            r=p.add_run();font(r,self.t,23,True,'accent');position_logo(r.add_picture(str(self.logo),width=Pt(w)))
        elif self.data.get('brand',{}).get('name'):font(p.add_run(text(self.data,self.data['brand']['name'],self.lang)),self.t,11,True)
        font(p.add_run('\t'+LABELS[self.lang]['specification']),self.t,8.5,color='muted')
        p=header.add_paragraph();p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=Pt(13)
        model=self.data['meta']['model']+' | '+LABELS[self.lang]['version']+' '+self.data['meta']['version']
        if self.data.get('fictional'):model+=' | '+LABELS[self.lang]['fiction']
        self.expected.append(model)
        font(p.add_run(model),self.t,8.5,color='muted')
        borders=el('pBdr');borders.append(el('bottom',val='single',sz=4,space=6,color=self.t['colors']['line']));p._element.get_or_add_pPr().append(borders)
        footer=s.footer
        for p in list(footer.paragraphs):p._element.getparent().remove(p._element)
        p=footer.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(11);p.paragraph_format.tab_stops.add_tab_stop(Pt(self.width),WD_TAB_ALIGNMENT.RIGHT)
        borders=el('pBdr');borders.append(el('top',val='single',sz=4,space=5,color=self.t['colors']['line']));p._element.get_or_add_pPr().append(borders)
        brand=self.data.get('brand',{});contact=' | '.join(brand[k] for k in ['website','email'] if brand.get(k));font(p.add_run(contact+'\t'),self.t,7.5,color='muted')
        r=p.add_run();font(r,self.t,10);r._element.append(el('fldChar',fldCharType='begin'));r=p.add_run();ins=el('instrText');ins.text=' PAGE ';r._element.append(ins);p.add_run()._element.append(el('fldChar',fldCharType='separate'));font(p.add_run('1'),self.t,10);p.add_run()._element.append(el('fldChar',fldCharType='end'))
        if brand.get('name'):
            p=footer.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(11);font(p.add_run(text(self.data,brand['name'],self.lang)),self.t,7.5,color='muted')

    def image(self,role,w,h,shade=False):
        img=self.data.get('images',{}).get(role)
        if not img:return
        path=asset_path(self.base,img['path']);im=Image.open(path);scale=min(w/im.width,h/im.height);iw,ih=im.width*scale,im.height*scale
        p=self.add('',after=8,align=WD_ALIGN_PARAGRAPH.CENTER,keep=True);p.paragraph_format.line_spacing=1
        if shade:p._element.get_or_add_pPr().append(el('shd',fill=self.t['colors']['pale']))
        p.add_run().add_picture(str(path),width=Pt(iw),height=Pt(ih))
        self.layout.append({'type':'image','role':role,'width_pt':iw,'height_pt':ih})
        if img.get('caption'):self.add(text(self.data,img['caption'],self.lang),size=9.5,leading=14.05,color='muted',after=9,align=WD_ALIGN_PARAGRAPH.CENTER)

    def table(self,headers,rows,widths,padding=3.35,row_ids=None):
        # Minimum heights are computed from the larger translation, while Word may
        # grow a row further. Exact-height clipping is deliberately prohibited.
        widths=[self.width*w/sum(widths) for w in widths];allrows=[headers]+rows;plans=[]
        for row in allrows:
            rh=max(height(text(self.data,value,lang),width-11,10,14.5,lang,False) for value,width in zip(row,widths) for lang in ('zh','en'))+padding*2
            require(rh<=575,'Specification record is taller than a usable page; split the record before generating')
            plans.append(max(21.2,rh))
        table=self.doc.add_table(rows=len(allrows),cols=len(widths));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        pr=table._tbl.tblPr
        for child in list(pr):
            if child.tag in [qn('w:tblBorders'),qn('w:tblCellMar'),qn('w:tblW')]:pr.remove(child)
        pr.append(el('tblW',type='dxa',w=round(self.width*20)))
        borders=el('tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:borders.append(el(side,val='single',sz=4,color=self.t['colors']['line']))
        pr.append(borders);m=el('tblCellMar')
        for side,val in [('top',round(padding*20)),('bottom',round(padding*20)),('left',110),('right',110)]:m.append(el(side,w=val,type='dxa'))
        pr.append(m)
        for col,width in zip(table.columns,widths):col.width=Pt(width)
        for index,(row,values,rh) in enumerate(zip(table.rows,allrows,plans)):
            # LibreOffice adds cell margins to its minimum row-content height.
            # Use the shared line-box minimum, not a padded height a second time.
            row.height=Pt(rh-padding*2);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST;tr=row._tr.get_or_add_trPr();tr.append(el('cantSplit'))
            if index==0:tr.append(el('tblHeader'))
            for cell,width,value in zip(row.cells,widths,values):
                cell.width=Pt(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tc=cell._tc.get_or_add_tcPr();tc.append(el('shd',fill=self.t['colors']['accent'] if index==0 else (self.t['colors']['pale'] if index%2 else 'FFFFFF')))
                p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(14.5);p.paragraph_format.keep_with_next=index==0;p.paragraph_format.keep_together=True
                s=text(self.data,value,self.lang);font(p.add_run(s),self.t,10,index==0);self.expected.append(s)
                if index==0:p.runs[0].font.color.rgb=RGBColor(255,255,255)
        self.layout.append({'type':'table','rows':len(rows),'columns':len(headers),'widths_pt':widths,'minimum_heights_pt':plans,'repeat_header':True,'split_rows':False,'row_ids':row_ids or []})
        return table

    def metrics(self):
        values=self.data.get('metrics',[])
        if not values:return
        headers=[m['label'] for m in values];cells=[]
        for m in values:
            s='{{fact:'+m['fact']+'}}';cells.append({l:s+(' '+text(self.data,m['unit_label'],l) if m.get('unit_label') else '') for l in ('zh','en')})
        table=self.table(headers,[cells],[1]*len(values),padding=6.5)
        for cell in table.rows[0].cells:
            cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in cell.paragraphs[0].runs:font(r,self.t,9.5,True);r.font.color.rgb=RGBColor(255,255,255)
        for cell in table.rows[1].cells:
            p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=Pt(25)
            for r in p.runs:font(r,self.t,19,True,'accent')


def build_specification(data,base,out,t,logo):
    result={}
    for lang in ('zh','en'):
        b=Spec(data,t,lang,base,logo);L=LABELS[lang];m=data['meta'];title=text(data,m['product'],lang)
        require(height(title,b.width,27,39.95,lang,True)<=119.85+.01,'Specification product title is too long; supply a concise formal product name')
        b.add(title,'Title');b.add(L['specification'].title() if lang=='en' else L['specification'],size=11,leading=16.3,color='muted',after=4,keep=True)
        b.add(m['model'],size=11.5,leading=17,bold=True,color='accent',after=15,keep=True)
        b.image('hero',b.width,198,True)
        b.add(L['overview'],'Heading 2');b.add(text(data,m['overview'],lang),after=13)
        b.metrics()
        if m.get('ordering'):b.add(L['ordering'],'Heading 2');b.add(text(data,m['ordering'],lang),after=0)
        b.add(L['hardware'],'Heading 1',page_break=True);b.add(m['model'],size=10.5,leading=15.55,bold=True,color='accent',after=8,keep=True)
        b.image('detail',b.width,145)
        for group in data['spec_groups']:
            b.add(text(data,group['title'],lang),'Heading 2')
            b.table([{'zh':LABELS['zh']['parameter'],'en':LABELS['en']['parameter']},{'zh':LABELS['zh']['value'],'en':LABELS['en']['value']}],[[r['label'],r['value']] for r in group['rows']],[2665,7801],row_ids=[r['id'] for r in group['rows']])
        if data.get('performance'):
            p=data['performance'];b.add(text(data,p['title'],lang),'Heading 1',page_break=True);b.add(text(data,p['intro'],lang),after=10)
            b.table(p['headers'],[r['cells'] for r in p['rows']],[1]*len(p['headers']),padding=4.1)
            b.add(text(data,p['conditions'],lang),size=9,leading=13.3,color='muted',before=7,after=0)
        if data.get('features') or data.get('applications') or data.get('brand'):
            b.add(L['features'],'Heading 1',page_break=True)
            if data.get('features'):
                b.add(L['key_features'],'Heading 2')
                for f in data['features']:b.add('• '+text(data,f['title'],lang)+': '+text(data,f['text'],lang),after=7)
            if data.get('applications'):
                b.add(L['applications'],'Heading 2');b.table([{'zh':LABELS['zh']['area'],'en':LABELS['en']['area']},{'zh':LABELS['zh']['example'],'en':LABELS['en']['example']}],[[a['title'],a['text']] for a in data['applications']],[2665,7801],padding=6)
            brand=data.get('brand',{});contacts=[]
            for key,label in [('name','company'),('phone','phone'),('website','website'),('email','email'),('address','address')]:
                if brand.get(key):contacts.append(L[label]+': '+text(data,brand[key],lang))
            if contacts:
                b.add(L['contact'],'Heading 2')
                for contact in contacts:b.add(contact,after=5,leading=14.8)
        for k in ['author','last_modified_by','subject','keywords','comments','category','identifier']:setattr(b.doc.core_properties,k,'')
        b.doc.core_properties.title=title;b.doc.core_properties.language='zh-CN' if lang=='zh' else 'en';b.doc.core_properties.created=datetime(2000,1,1,tzinfo=timezone.utc);b.doc.core_properties.modified=b.doc.core_properties.created
        b.doc.settings.element.append(el('updateFields',val='true'))
        path=Path(out)/f'specification-{lang}.docx';b.doc.save(path)
        result[lang]={'docx':path.name,'expected_text':b.expected,'layout':b.layout,'logo_visible_width_pt':t['logo_visible_width_mm']*72/25.4 if logo else None}
    return result
