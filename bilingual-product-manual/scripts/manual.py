#!/usr/bin/env python3
"""Create paired manuals from one evidence-backed bilingual data model."""
from __future__ import annotations
import argparse, json, re, sys, math
from copy import deepcopy
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
LANGS = ('zh', 'en')
FACT = re.compile(r'\{\{fact:([a-z][a-z0-9_-]*)\}\}')
PLACEHOLDER = re.compile(r'\{\{|\}\}|\[\[|\]\]|<[^>]*(?:placeholder|insert)[^>]*>|\b(?:TODO|TBD|LOREM IPSUM)\b', re.I)
LABEL = {'zh': {'kicker':'软件文档', 'manual':'用户操作手册', 'version':'版本', 'applies':'适用软件版本', 'contents':'目录', 'figure':'图', 'fiction':'虚构产品示例 · 仅供模板演示'},
         'en': {'kicker':'USER MANUAL', 'manual':'User Manual', 'version':'Version', 'applies':'For software version', 'contents':'Contents', 'figure':'Figure', 'fiction':'FICTIONAL PRODUCT · TEMPLATE DEMONSTRATION ONLY'}}

class InputError(ValueError): pass

def require(ok, message):
    if not ok: raise InputError(message)

def keys(obj, allowed, required, where):
    require(isinstance(obj, dict), f'{where}: expected object')
    require(not (set(obj)-set(allowed)), f'{where}: unknown fields {sorted(set(obj)-set(allowed))}')
    require(set(required)<=set(obj), f'{where}: missing fields {sorted(set(required)-set(obj))}')

def pair(v, where, facts, check_numbers=True):
    keys(v, LANGS, LANGS, where)
    for lang in LANGS:
        require(isinstance(v[lang],str) and v[lang].strip(), f'{where}.{lang}: nonempty translation required')
        require(len(v[lang])<=12000, f'{where}.{lang}: split long content into smaller blocks')
        expand(v[lang],facts)
    if check_numbers:
        numbers=lambda s: sorted(re.findall(r'[-+]?\d+(?:\.\d+)*', s.replace(',', '').replace('−','-')))
        require(numbers(expand(v['zh'],facts))==numbers(expand(v['en'],facts)), f'{where}: bilingual numeric mismatch; use shared facts for values and units')
        units=lambda s: sorted(re.findall(r'[-+]?\d+(?:\.\d+)?\s*(ms|s|min|h|mm|cm|km|m|mA|A|mV|V|kW|W|kWh|Hz|kHz|MHz|GHz|MB|GB|TB|Mbps|Gbps|fps|°C|%)(?![A-Za-z])',s))
        require(units(expand(v['zh'],facts))==units(expand(v['en'],facts)), f'{where}: bilingual unit mismatch; use shared fact tokens')

def expand(text, facts):
    def sub(m):
        require(m[1] in facts, f'Unknown fact {m[1]}')
        f=facts[m[1]]
        return f['value'] + ((' '+f['unit']) if f.get('unit') else '')
    text=FACT.sub(sub,text)
    require(not PLACEHOLDER.search(text), 'Unresolved template marker or draft token in input')
    return text

def asset_path(base, value):
    require(isinstance(value,str) and value.strip(), 'Asset path must be a nonempty relative path')
    rel=Path(value)
    require(not rel.is_absolute() and '..' not in rel.parts, f'Asset must stay inside the input directory: {value}')
    path=(base/rel).resolve()
    require(path.is_relative_to(base.resolve()) and path.is_file(), f'Asset missing or outside input directory: {value}')
    require(path.suffix.lower() in {'.png','.jpg','.jpeg'}, f'Use PNG or JPEG: {value}')
    return path

def validate(data, base):
    keys(data, ['schema_version','template','fictional','meta','brand','facts','sources','sections','gaps','theme'], ['schema_version','template','meta','sources','sections'], 'input')
    require(data['schema_version']==1, 'schema_version must be 1')
    require(data['template']=='software-manual', 'Only software-manual is supported; do not mix in brochure/specification layouts')
    require(isinstance(data.get('fictional',False),bool),'fictional must be boolean')
    facts=data.get('facts',{})
    require(isinstance(facts,dict),'facts must be an object')
    sources=data['sources'];require(isinstance(sources,list) and len(sources)>0,'At least one source is required')
    source_ids=set()
    for x in sources:
        keys(x,['id','description','fictional'],['id','description'],'source')
        require(isinstance(x['id'],str) and x['id'] not in source_ids,'Source IDs must be unique strings')
        require(isinstance(x['description'],str) and x['description'].strip(),'Source description required')
        source_ids.add(x['id'])
        require(not x.get('fictional') or data.get('fictional'), 'Fictional sources require a visible fictional-product label')
    def refs(obj,where):
        r=obj.get('sources',[])
        require(isinstance(r,list) and r and all(isinstance(x,str) and x in source_ids for x in r),f'{where}: valid evidence source IDs required')
    for key,f in facts.items():
        require(re.fullmatch('[a-z][a-z0-9_-]*',key),f'Invalid fact ID: {key}')
        keys(f,['value','unit','sources'],['value','sources'],f'fact {key}')
        require(isinstance(f['value'],str) and f['value'].strip(),f'Fact {key}: value must be a string')
        require(isinstance(f.get('unit',''),str),f'Fact {key}: unit must be a string')
        require(not PLACEHOLDER.search(f['value']+f.get('unit','')),f'Fact {key}: unresolved marker')
        refs(f,f'fact {key}')
    meta=data['meta']
    keys(meta,['product','title','short_title','version','applies_to','intro'],['product','version','intro'],'meta')
    for k in ['product','title','short_title','applies_to','intro']:
        if k in meta:pair(meta[k],'meta.'+k,facts)
    require(isinstance(meta['version'],str) and meta['version'].strip() and len(meta['version'])<60,'version must be a short shared string')
    expand(meta['version'],facts)
    for lang in LANGS:
        require(len(meta.get('short_title',meta['product'])[lang])<=90,'Provide a short_title of at most 90 characters for running headers')
    brand=data.get('brand',{})
    keys(brand,['name','contact','logo'],[],'brand')
    for k in ['name','contact']:
        if k in brand:pair(brand[k],'brand.'+k,facts)
    if brand.get('logo'):asset_path(base,brand['logo'])
    theme=data.get('theme',{})
    keys(theme,['fonts','text_color','caption_color','table_header_fill','table_border','toc_depth'],[],'theme')
    if 'fonts' in theme:
        keys(theme['fonts'],['latin','cjk','code'],[],'theme.fonts')
        require(all(isinstance(v,str) and v.strip() for v in theme['fonts'].values()),'Font names must be strings')
    for k in ['text_color','caption_color','table_header_fill','table_border']:
        if k in theme:require(isinstance(theme[k],str) and re.fullmatch('[0-9A-Fa-f]{6}',theme[k]), f'{k}: expected six hex digits')
    require(type(theme.get('toc_depth',2)) is int and theme.get('toc_depth',2) in (1,2),'toc_depth must be 1 or 2')
    sections=data['sections'];require(isinstance(sections,list) and sections,'At least one section is required')
    seen=set();previous=0;fig_ids=set()
    for s in sections:
        keys(s,['id','level','title','blocks','page_break'],['id','level','title','blocks'],'section')
        require(isinstance(s['id'],str) and re.fullmatch('[a-z][a-z0-9_-]{0,30}',s['id']) and s['id'] not in seen,'Section IDs must be unique lowercase identifiers')
        seen.add(s['id']);level=s['level']
        require(type(level) is int and 1<=level<=4 and level<=previous+1,'Heading levels must start at 1, stay within 1..4 and never skip a level')
        previous=level;pair(s['title'],'section '+s['id'],facts)
        require(isinstance(s['blocks'],list),'blocks must be a list')
        if 'page_break' in s:require(isinstance(s['page_break'],bool),'page_break must be boolean')
        for b in s['blocks']:
            require(isinstance(b,dict),'Block must be an object')
            t=b.get('type')
            if t in ['paragraph','step','note','code']:
                keys(b,['type','text','sources'],['type','text','sources'],t)
                pair(b['text'],t,facts);refs(b,t)
            elif t=='figure':
                keys(b,['type','id','path','caption','alt','width_in','sources'],['type','id','path','caption','alt','sources'],'figure')
                require(isinstance(b['id'],str) and b['id'] not in fig_ids,'Figure IDs must be unique strings');fig_ids.add(b['id'])
                path=b['path']
                if isinstance(path,dict):
                    keys(path,LANGS,LANGS,'figure.path')
                    for v in path.values():asset_path(base,v)
                else:asset_path(base,path)
                pair(b['caption'],'figure.caption',facts);pair(b['alt'],'figure.alt',facts);refs(b,'figure')
                require(isinstance(b.get('width_in',6.9),(int,float)) and .4<=b.get('width_in',6.9)<=6.9,'figure width_in must be 0.4..6.9')
            elif t=='table':
                keys(b,['type','caption','headers','rows','widths','sources'],['type','headers','rows','sources'],'table')
                require(isinstance(b['headers'],list) and 2<=len(b['headers'])<=6,'Tables require 2..6 columns')
                if 'caption' in b:pair(b['caption'],'table.caption',facts)
                for h in b['headers']:pair(h,'table.header',facts)
                require(isinstance(b['rows'],list) and b['rows'],'Table needs data rows')
                for row in b['rows']:
                    require(isinstance(row,list) and len(row)==len(b['headers']),'Table row/column mismatch')
                    for cell in row:
                        pair(cell,'table.cell',facts)
                        require(max(map(len,cell.values()))<=1200,'Split very long table cells into smaller records or prose')
                if 'widths' in b:require(isinstance(b['widths'],list) and len(b['widths'])==len(b['headers']) and all(isinstance(x,(int,float)) and x>0 for x in b['widths']),'widths must be positive column weights')
                weights=b.get('widths',[1]*len(b['headers']))
                for row in [b['headers']]+b['rows']:
                    for cell,weight in zip(row,weights):
                        width=6.9*72*weight/sum(weights)-9
                        require(width>=24,'Table column is too narrow; rebalance widths')
                        for text in cell.values():
                            estimated=sum(1 if '\u2e80'<=c<='\uffff' else .55 for c in expand(text,facts))*10/width
                            require((math.ceil(estimated)+text.count('\n'))*13+9<=530,'Table record is taller than a page; split this record or move its explanation into prose')
                refs(b,'table')
            else:raise InputError(f'Unsupported block type: {t}')
    gaps=data.get('gaps',[]);require(isinstance(gaps,list),'gaps must be a list')
    for g in gaps:
        keys(g,['field','reason','action'],['field','reason','action'],'gap')
        require(all(isinstance(g[x],str) and g[x].strip() for x in g),'Gap fields must be nonempty strings')
        require(g['action'] in ('omit','internal'),'Gap action must be omit or internal')
    return facts

def load(path):
    path=Path(path).resolve();data=json.loads(path.read_text(encoding='utf-8'));validate(data,path.parent);return data,path.parent

def theme_for(data):
    t=json.loads((ROOT/'assets/theme.json').read_text());override=data.get('theme',{})
    t['fonts'].update(override.get('fonts',{}));t.update({k:v for k,v in override.items() if k!='fonts'});return t

def paragraph(doc,text,style='Normal'):
    role=next(s for s in doc.styles if s.style_id==style)
    p=doc.add_paragraph(style=role)
    for i,s in enumerate(re.split(r'\*\*(.*?)\*\*',text,flags=re.S)):
        r=p.add_run(s)
        if i%2:r.bold=True
    p.paragraph_format.widow_control=True
    return p

def field(p, instruction, cache='1'):
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
    r=p.add_run();a=OxmlElement('w:instrText');a.set(qn('xml:space'),'preserve');a.text=' '+instruction+' ';r._r.append(a)
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'separate');r._r.append(a)
    p.add_run(cache)
    r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');r._r.append(a)

def bookmark(p,name,index):
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(index));a.set(qn('w:name'),name)
    p._p.insert(1,a);a=OxmlElement('w:bookmarkEnd');a.set(qn('w:id'),str(index));p._p.append(a)

def heading_inventory(data,lang):
    nums=[0,0,0,0];out=[]
    for i,s in enumerate(data['sections']):
        lv=s['level'];nums[lv-1]+=1;nums[lv:]=[0]*(4-lv)
        title='.'.join(str(x) for x in nums[:lv])+' '+expand(s['title'][lang],data.get('facts',{}))
        out.append({'id':s['id'],'level':lv,'title':title,'bookmark':'sec_'+s['id'].replace('-','_')})
    return out

def set_style_fonts(d,lang,theme):
    for st in d.styles:
        if st.type!=1:continue
        rp=st.element.get_or_add_rPr();f=rp.find(qn('w:rFonts'))
        if f is None:f=OxmlElement('w:rFonts');rp.insert(0,f)
        latin=theme['fonts']['code'] if st.style_id=='Code' else theme['fonts']['latin']
        for k in ['ascii','hAnsi','cs']:f.set(qn('w:'+k),latin)
        f.set(qn('w:eastAsia'),theme['fonts']['cjk'] if lang=='zh' else latin)
        st.font.color.rgb=RGBColor.from_string(theme['caption_color'] if st.style_id in ('Caption','Footer') else theme['text_color'])

def add_table(d,b,lang,facts,theme):
    if b.get('caption'):
        p=paragraph(d,expand(b['caption'][lang],facts),'Caption');p.paragraph_format.keep_with_next=True
    n=len(b['headers']);table=d.add_table(rows=1,cols=n);table.autofit=False
    widths=b.get('widths',[1]*n);total=sum(widths);widths=[6.9*x/total for x in widths]
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),theme['table_border']);borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)
    for col,w in zip(table.columns,widths):col.width=Inches(w)
    for ri,rowdata in enumerate([b['headers']]+b['rows']):
        row=table.rows[0] if ri==0 else table.add_row()
        trpr=row._tr.get_or_add_trPr()
        if ri==0:
            rep=OxmlElement('w:tblHeader');trpr.append(rep)
        # Keep a complete record together; input validation rejects records
        # too tall to fit one page so cantSplit cannot cause silent clipping.
        trpr.append(OxmlElement('w:cantSplit'))
        for ci,v in enumerate(rowdata):
            cell=row.cells[ci];cell.width=Inches(widths[ci]);cell.vertical_alignment=1
            p=cell.paragraphs[0];p.style='Normal';p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.keep_with_next=(ri==0)
            text=expand(v[lang],facts)
            r=p.add_run(text);r.font.size=Pt(10);r.bold=(ri==0)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),theme['table_header_fill']);cell._tc.get_or_add_tcPr().append(sh)
    paragraph(d,'').paragraph_format.space_after=Pt(3)

def build_one(data,base,lang,dest):
    t=theme_for(data);facts=data.get('facts',{});labels=LABEL[lang];m=data['meta'];brand=data.get('brand',{})
    d=Document(ROOT/'assets'/f'manual-{lang}.docx');set_style_fonts(d,lang,t)
    title=expand(m.get('title',m['product'])[lang],facts)
    if 'title' not in m:title+='\n'+labels['manual']
    estimate=sum(1 if '\u2e80'<=c<='\uffff' else .52 for c in title)*29/(6.9*72)+title.count('\n')
    require(estimate<11,'Cover title is too long; supply a shorter meta.title while preserving the full product name in body content')
    sec=d.sections[0];sec.different_first_page_header_footer=True
    if brand.get('logo'):
        sec.first_page_header.paragraphs[0].add_run().add_picture(str(asset_path(base,brand['logo'])),width=Inches(1.1))
    hp=sec.header.paragraphs[0];hp.style=next(s for s in d.styles if s.style_id=='Header');hp.add_run(expand(m.get('short_title',m['product'])[lang],facts))
    fp=sec.footer.paragraphs[0];fp.style=next(s for s in d.styles if s.style_id=='Footer')
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT)
    fp.add_run(labels['manual']+'  |  '+m['version']+'\t');field(fp,'PAGE');fp.add_run(' / ');field(fp,'NUMPAGES')
    p=paragraph(d,labels['kicker']);p.paragraph_format.space_before=Pt(t['cover_top_space_pt']);p.paragraph_format.space_after=Pt(18)
    for r in p.runs:r.bold=True;r.font.size=Pt(10)
    paragraph(d,title,'Title')
    paragraph(d,labels['version']+' '+m['version'],'Subtitle')
    if m.get('applies_to'):paragraph(d,expand(m['applies_to'][lang],facts),'CoverMetadata')
    else:paragraph(d,labels['applies']+' '+m['version'],'CoverMetadata')
    if data.get('fictional'):
        p=paragraph(d,labels['fiction'],'CoverMetadata');p.paragraph_format.space_before=Pt(14)
        for r in p.runs:r.font.size=Pt(9)
    if brand.get('name'):
        p=paragraph(d,expand(brand['name'][lang],facts),'CoverMetadata')
        p.paragraph_format.space_before=Pt(max(24,t['cover_brand_space_pt']-max(0,math.ceil(estimate)-2)*32-(25 if data.get('fictional') else 0)))
        for r in p.runs:r.bold=True
    if brand.get('contact'):paragraph(d,expand(brand['contact'][lang],facts),'CoverMetadata')
    d.add_page_break()
    p=paragraph(d,labels['contents'],'Title');p.paragraph_format.space_after=Pt(14)
    p=paragraph(d,expand(m['intro'][lang],facts));p.paragraph_format.space_after=Pt(14)
    inventory=heading_inventory(data,lang)
    for h in inventory:
        if h['level']>t['toc_depth']:continue
        p=paragraph(d,h['title'],'Contents'+str(h['level']));p.paragraph_format.keep_with_next=False
        p.paragraph_format.right_indent=Inches(.18)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        p.add_run('\t');field(p,'PAGEREF '+h['bookmark']+' \\h','1')
    d.add_page_break();figure_no=0
    for index,(s,h) in enumerate(zip(data['sections'],inventory),1):
        p=paragraph(d,h['title'],'Heading'+str(s['level']));bookmark(p,h['bookmark'],index)
        # Original short chapters 1/2 share a page; explicit section boundaries
        # carry the source's later chapter-start behavior without blank pages.
        p.paragraph_format.page_break_before=s.get('page_break',False) if index>1 else False
        step_no=0
        for b in s['blocks']:
            typ=b['type']
            if typ in ('paragraph','note','step','code'):
                text=expand(b['text'][lang],facts)
                if typ=='step':step_no+=1;text=f'{step_no}. '+text
                p=paragraph(d,text,{'step':'BodyStep','code':'Code'}.get(typ,'Normal'))
                if typ=='note':
                    for r in p.runs:r.italic=True
            elif typ=='table':add_table(d,b,lang,facts,t)
            elif typ=='figure':
                figure_no+=1;path=asset_path(base,b['path'][lang] if isinstance(b['path'],dict) else b['path'])
                with Image.open(path) as im:w,hp=im.size
                width=min(b.get('width_in',t['figure_max_width_in']),t['figure_max_width_in'],t['figure_max_height_in']*w/hp)
                p=paragraph(d,'','Picture');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
                sh=p.add_run().add_picture(str(path),width=Inches(width));sh._inline.docPr.set('descr',expand(b['alt'][lang],facts))
                p=paragraph(d,labels['figure']+' '+str(figure_no)+' '+expand(b['caption'][lang],facts),'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
    st=d.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');st.append(u)
    for prop in ['author','last_modified_by','subject','keywords','comments','category','identifier']:
        setattr(d.core_properties,prop,'')
    d.core_properties.title=title.replace('\n',' ');d.core_properties.language='zh-CN' if lang=='zh' else 'en-GB'
    d.core_properties.created=datetime(2000,1,1,tzinfo=timezone.utc);d.core_properties.modified=d.core_properties.created
    d.save(dest)
    return inventory

def build(path,out):
    data,base=load(path);out=Path(out).resolve();out.mkdir(parents=True,exist_ok=True)
    require(not any((out/f'manual-{l}.docx').exists() for l in LANGS),'Output exists; use a fresh output directory')
    manifest={'schema_version':1,'template':data['template'],'fictional':data.get('fictional',False),'languages':{},'facts':data.get('facts',{}),'gaps':data.get('gaps',[]),'fonts':theme_for(data)['fonts'],'status':'docx-built-render-required'}
    try:
        for lang in LANGS:
            p=out/f'manual-{lang}.docx';inv=build_one(data,base,lang,p)
            manifest['languages'][lang]={'docx':p.name,'headings':inv}
    except Exception:
        for lang in LANGS:(out/f'manual-{lang}.docx').unlink(missing_ok=True)
        raise
    (out/'build-report.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    # Gaps remain internal and are never injected as public manual prose.
    print(json.dumps({'result':'docx-built','output':str(out),'internal_gaps':len(manifest['gaps']),'next':'render.py then inspect every PNG'},ensure_ascii=False))

def main():
    ap=argparse.ArgumentParser(description=__doc__);sub=ap.add_subparsers(dest='command',required=True)
    p=sub.add_parser('validate');p.add_argument('input')
    p=sub.add_parser('build');p.add_argument('input');p.add_argument('--out',required=True)
    args=ap.parse_args()
    try:
        if args.command=='validate':load(args.input);print('Input validation passed')
        else:build(args.input,args.out)
    except (InputError,KeyError,TypeError,json.JSONDecodeError) as exc:
        print(f'INPUT ERROR: {exc}',file=sys.stderr);return 2
    return 0
if __name__=='__main__':sys.exit(main())
